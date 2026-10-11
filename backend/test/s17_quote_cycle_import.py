# -*- coding: utf-8 -*-
"""
S17 每日寄语批量导入与按天循环 冒烟脚本（TestClient + SQLite 临时库，不依赖 8010/openGauss）

覆盖 V1.1 功能二验收点：
  C01 空数据：无启用寄语时 /quotes/today 返回 200 null，循环概览 enabled_count=0
  C02 模板下载：xlsx 双工作表、表头正确
  C03 合法文件预览 + 确认导入：统计正确，节点名/ID 解析，停用状态，空排序号自动接续
  C04 同一 token 重复确认 → 400
  C05 错误行定位：内容为空/排序号非法/启用状态非法/节点不存在/库内重复/文件内重复/公式注入；整批确认 400
  C06 非 .xlsx 与表头不符 → 400
  C07 按天循环：起始日=今天→第1条，起始日=昨天→第2条，起始日=今天-N→回到第1条，起始日在未来→兜底第1条
  C08 同一天不同用户/重复请求结果一致，date 为当天
  C09 停用寄语不参与循环；启用条数随之变化
  C10 排序号相同按 id 升序（稳定次序）；修改排序号后按新顺序计算
  C11 删除参与循环的寄语后不越界；只剩一条时每天展示该条
  C12 管理端校验：内容空/排序号越界/节点不存在 400；编辑不存在 404；起始日期非法 400
  C13 鉴权：/quotes/today 无 token 401；/admin/quotes/cycle 用户 token 401/403

用法：cd backend && python test/s17_quote_cycle_import.py
"""
import io
import json
import os
import sys
from datetime import date, timedelta

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

_DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_smoke_quote.db")
if os.path.exists(_DB_PATH):
    os.remove(_DB_PATH)
os.environ["DATABASE_URL"] = "sqlite:///" + _DB_PATH.replace("\\", "/")

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.database import Base, SessionLocal, engine
from app.core.helpers import today_str
from app.core.security import create_access_token, create_admin_token
from app.data.seed import init_seed
from app.main import app
from app.models.models import DailyQuote, User

RESULTS = []
HEADERS = ["寄语内容", "出处", "排序号", "启用状态", "关联节点"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TODAY = date.fromisoformat(today_str())


def record(case_id, name, ok, detail=""):
    RESULTS.append({"id": case_id, "name": name, "ok": bool(ok), "detail": str(detail)[:300]})
    print(("PASS " if ok else "FAIL ") + case_id + " " + name + ((" | " + str(detail)[:200]) if detail else ""))


def build_xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_user(openid):
    db = SessionLocal()
    try:
        user = User(openid=openid, nickname=openid)
        db.add(user)
        db.commit()
        return {"Authorization": "Bearer " + create_access_token(user.id, openid)}
    finally:
        db.close()


def main():
    Base.metadata.create_all(engine)
    db = SessionLocal()
    try:
        init_seed(db)
        db.query(DailyQuote).delete()
        db.commit()
    finally:
        db.close()

    client = TestClient(app)
    admin = {"Authorization": "Bearer " + create_admin_token("s17-smoke")}
    user_a = make_user("s17-user-a")
    user_b = make_user("s17-user-b")

    def today_quote(headers=user_a):
        r = client.get("/api/quotes/today", headers=headers)
        return r.status_code, r.json()

    def set_start(d):
        return client.put("/api/admin/quotes/cycle", headers=admin, json={"start_date": d.isoformat()})

    def preview(content, name="quotes.xlsx"):
        return client.post("/api/admin/quotes/import-preview", headers=admin, files={"file": (name, content, XLSX)})

    try:
        # C01 空数据
        st, body = today_quote()
        cyc = client.get("/api/admin/quotes/cycle", headers=admin).json()
        record("C01", "无启用寄语：today 200 null，enabled_count=0",
               st == 200 and body is None and cyc["enabled_count"] == 0 and cyc["today_quote_id"] is None,
               f"status={st} body={body} cycle={cyc}")

        # C02 模板下载
        r = client.get("/api/admin/quotes/import-template", headers=admin)
        ok, detail = False, f"status={r.status_code}"
        if r.status_code == 200:
            wb = load_workbook(io.BytesIO(r.content))
            got = [c.value for c in wb["寄语数据"][1]]
            ok = "填写说明" in wb.sheetnames and got == HEADERS and "attachment" in r.headers.get("content-disposition", "")
            detail = f"sheets={wb.sheetnames} header={got}"
        record("C02", "模板下载：xlsx 双工作表且表头正确", ok, detail)

        # C03 合法导入：排序号 1/留空/3/留空，第 3 行停用，节点名与 ID
        rows = [
            ["S17寄语一", "出处一", 1, "启用", "瑞金"],
            ["S17寄语二", "出处二", "", "", 3],
            ["S17寄语三", "出处三", 3, "停用", ""],
            ["S17寄语四", "出处四", "", "是", ""],
        ]
        r = preview(build_xlsx(HEADERS, rows))
        body = r.json() if r.status_code == 200 else {}
        token = body.get("preview_token", "")
        ok_preview = r.status_code == 200 and body.get("valid_count") == 4 and body.get("invalid_count") == 0
        r = client.post("/api/admin/quotes/import-confirm", headers=admin, json={"preview_token": token})
        lst = client.get("/api/admin/quotes?page=1&page_size=100", headers=admin).json()
        by_content = {i["content"]: i for i in lst["items"]}
        q1, q2, q3, q4 = (by_content.get(f"S17寄语{n}") for n in "一二三四")
        ok_data = all([q1, q2, q3, q4]) and (
            q1["node_name"] == "瑞金" and q2["node_id"] == 3 and q3["is_enabled"] is False
            and q1["sort_order"] == 1 and q3["sort_order"] == 3
            and q2["sort_order"] == 4 and q4["sort_order"] == 5
        )
        record("C03", "合法导入：统计、节点名/ID、停用、空排序号接续最大值之后",
               ok_preview and r.status_code == 200 and r.json().get("imported") == 4 and ok_data,
               f"preview={ok_preview} confirm={r.status_code} orders="
               f"{[(i['content'], i['sort_order'], i['is_enabled'], i['node_id']) for i in lst['items']]}")

        # C04 重复确认
        r = client.post("/api/admin/quotes/import-confirm", headers=admin, json={"preview_token": token})
        record("C04", "同一 token 重复确认被拒绝", r.status_code == 400, f"status={r.status_code} detail={r.json().get('detail')}")

        # C05 错误行定位
        bad_rows = [
            ["", "出处", 1, "", ""],                    # row2 内容为空
            ["S17坏排序", "出处", "1.5", "", ""],       # row3 排序号非法
            ["S17坏状态", "出处", 1, "未知", ""],       # row4 启用状态非法
            ["S17坏节点", "出处", 1, "", "不存在的节点"],  # row5 节点不存在
            ["S17寄语一", "出处", 1, "", ""],           # row6 与库内重复
            ["S17文件内重复", "出处", 1, "", ""],       # row7
            ["S17文件内重复", "出处", 2, "", ""],       # row8 文件内重复
            ["=HYPERLINK(\"x\")", "出处", 1, "", ""],   # row9 公式注入
        ]
        r = preview(build_xlsx(HEADERS, bad_rows))
        body = r.json() if r.status_code == 200 else {}
        err_rows = sorted(e["row"] for e in body.get("errors", []))
        err_fields = {e["row"]: e["field"] for e in body.get("errors", [])}
        expected_fields = {2: "寄语内容", 3: "排序号", 4: "启用状态", 5: "关联节点", 6: "寄语内容", 8: "寄语内容"}
        r2 = client.post("/api/admin/quotes/import-confirm", headers=admin, json={"preview_token": body.get("preview_token", "")})
        record("C05", "错误行定位到行与字段且含错批次禁止确认",
               r.status_code == 200 and err_rows == [2, 3, 4, 5, 6, 8, 9] and body.get("valid_count") == 1
               and all(err_fields.get(k) == v for k, v in expected_fields.items()) and r2.status_code == 400,
               f"rows={err_rows} valid={body.get('valid_count')} confirm={r2.status_code} "
               f"errors={json.dumps(body.get('errors', []), ensure_ascii=False)[:200]}")

        # C06 非 xlsx / 表头不符
        r1 = preview(b"a,b", "quotes.csv")
        r2 = preview(build_xlsx(["日期", "寄语内容", "出处"], [["2026-10-01", "x", "y"]]))
        record("C06", "非 .xlsx 与表头不符被拒绝",
               r1.status_code == 400 and r2.status_code == 400 and "表头" in r2.json().get("detail", ""),
               f"csv={r1.status_code} header={r2.status_code}")

        # 启用列表（按 sort_order, id）：寄语一(1) 寄语二(4) 寄语四(5)，N=3
        order = [q1["id"], q2["id"], q4["id"]]

        # C07 按天循环
        got = {}
        for label, start in (("d0", TODAY), ("d1", TODAY - timedelta(days=1)), ("d2", TODAY - timedelta(days=2)),
                             ("wrap", TODAY - timedelta(days=3)), ("future", TODAY + timedelta(days=5))):
            set_start(start)
            got[label] = today_quote()[1]["id"]
        cyc = client.get("/api/admin/quotes/cycle", headers=admin).json()
        record("C07", "按天循环：第1/2/3条、末条后回到第1条、起始日在未来兜底第1条",
               got == {"d0": order[0], "d1": order[1], "d2": order[2], "wrap": order[0], "future": order[0]}
               and cyc["enabled_count"] == 3 and cyc["today_position"] == 1,
               f"got={got} expected={order} cycle={cyc}")

        # C08 同一天一致
        set_start(TODAY - timedelta(days=1))
        a1, a2, b1 = today_quote(user_a)[1], today_quote(user_a)[1], today_quote(user_b)[1]
        record("C08", "同一天不同用户/重复请求得到同一条，date 为当天，含出处与节点",
               a1["id"] == a2["id"] == b1["id"] == order[1] and a1["date"] == TODAY.isoformat()
               and a1["source"] == "出处二" and a1["node"] and a1["node"]["id"] == 3,
               f"a1={a1} b1={b1['id']}")

        def put_quote(q, **changes):
            payload = {k: q[k] for k in ("content", "source", "sort_order", "is_enabled", "node_id")}
            payload.update(changes)
            return client.put(f"/api/admin/quotes/{q['id']}", headers=admin, json=payload)

        # C09 停用不参与循环：停用寄语二 → 列表 [一, 四]，起始日=昨天 → 四
        put_quote(q2, is_enabled=False)
        cyc = client.get("/api/admin/quotes/cycle", headers=admin).json()
        record("C09", "停用寄语不参与循环，按最新启用列表重算",
               today_quote()[1]["id"] == q4["id"] and cyc["enabled_count"] == 2,
               f"today={today_quote()[1]['id']} expected={q4['id']} cycle={cyc}")

        # C10 排序号相同按 id 升序；修改排序号后按新顺序
        put_quote(q4, sort_order=1)          # 一(1,id小) 四(1,id大)
        set_start(TODAY)
        tie_first = today_quote()[1]["id"]
        put_quote(q4, sort_order=0)          # 四(0) 一(1)
        new_first = today_quote()[1]["id"]
        record("C10", "排序号相同按 id 升序；修改排序号后按新顺序",
               tie_first == q1["id"] and new_first == q4["id"],
               f"tie_first={tie_first} expected={q1['id']} new_first={new_first} expected={q4['id']}")

        # C11 删除参与循环的寄语：启用 [四, 一]，起始日=昨天 → 一；删除一后只剩四，任意起始日都是四
        set_start(TODAY - timedelta(days=1))
        before = today_quote()[1]["id"]
        client.delete(f"/api/admin/quotes/{q1['id']}", headers=admin)
        after = [None, None]
        for i, start in enumerate((TODAY - timedelta(days=1), TODAY - timedelta(days=7))):
            set_start(start)
            st, body = today_quote()
            after[i] = body["id"] if st == 200 and body else (st, body)
        record("C11", "删除循环中寄语不越界，仅剩一条时每天展示该条",
               before == q1["id"] and after == [q4["id"], q4["id"]],
               f"before={before} after={after} expected={q4['id']}")

        # C12 管理端校验
        base = {"content": "S17手工", "source": "出处", "sort_order": 1, "is_enabled": True, "node_id": None}
        s_empty = client.post("/api/admin/quotes", headers=admin, json=dict(base, content=" ")).status_code
        s_sort = client.post("/api/admin/quotes", headers=admin, json=dict(base, sort_order=-1)).status_code
        s_node = client.post("/api/admin/quotes", headers=admin, json=dict(base, node_id=9999)).status_code
        s_404 = client.put("/api/admin/quotes/999999", headers=admin, json=base).status_code
        s_date = client.put("/api/admin/quotes/cycle", headers=admin, json={"start_date": "2026-02-30"}).status_code
        created = client.post("/api/admin/quotes", headers=admin, json=dict(base, node_id=1))
        record("C12", "管理端校验 400/404 与正常新增",
               (s_empty, s_sort, s_node, s_404, s_date) == (400, 400, 400, 404, 400)
               and created.status_code == 200 and created.json()["node_name"] == "瑞金",
               f"empty={s_empty} sort={s_sort} node={s_node} 404={s_404} date={s_date} create={created.status_code}")

        # C13 鉴权
        s1 = client.get("/api/quotes/today").status_code
        s2 = client.get("/api/admin/quotes/cycle", headers=user_a).status_code
        record("C13", "鉴权：today 无 token 401；用户 token 访问管理端 401/403",
               s1 == 401 and s2 in (401, 403), f"today={s1} admin_by_user={s2}")
    finally:
        engine.dispose()

    passed = sum(1 for x in RESULTS if x["ok"])
    print(f"\n== S17 每日寄语导入与循环: {passed}/{len(RESULTS)} 通过 ==")
    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "s17-quote-cycle-results.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump({"results": RESULTS}, f, ensure_ascii=False, indent=2)
    if os.path.exists(_DB_PATH):
        os.remove(_DB_PATH)
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    sys.exit(main())
