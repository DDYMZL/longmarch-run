# -*- coding: utf-8 -*-
"""
S25 AI 智能出题 冒烟脚本（TestClient + SQLite 临时库，不依赖 8010/openGauss）

覆盖需求验收点：
  A01 Excel 模板下载未受影响（回归）
  A02 Excel 预览 + 确认导入未受影响（回归）
  A03 生成接口参数校验：空主题 / 题型非单选 / 难度非法 / 分类非法 / 数量越界
  A04 鉴权：未登录 401、无 questions 菜单的管理员 403（生成与保存均拦）
  A05 模型故障映射：超时 / 鉴权失败 / 服务异常 / 回复截断 / 无效 JSON
  A06 输出容错：Markdown 代码围栏与前后多余文本仍能解析
  A07 逐题问题识别：答案不在选项 / 选项不足 / 题干为空 / 与库内重复 / 批内重复
  A08 生成结果为草稿，不写入正式题库
  A09 批量保存重新校验：合法项入库且 id 连续、非法项计入 failures（部分失败）
  A10 批量保存重复题拦截（库内重复 + 批内重复）
  A11 批量保存空列表 400
  A12 保存后题库列表可查到，字段与 Excel 导入题目同构
  A13 任何响应都不泄露 LLM_API_KEY
  A14 真实调用模型：生成 10 道单选题，结构与答案格式合规
  A15 真实调用模型：生成 50 道题，分批处理长响应无批次失败
  A16 小程序答题契约不受影响：下发 camelCase 且不含 answer/analysis

用法：
  cd backend
  # 本机 DNS 不可用时，用 DoH 查到的 IP 建立测试专用映射（不影响生产代码）
  set LM_TEST_DNS_MAP=dashscope.aliyuncs.com=<ip>
  python test/s25_ai_question_generate.py
  # 跳过耗时的 50 题真实生成：set LM_AI_SKIP_LARGE=1
"""
import io
import json
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

_HERE = os.path.dirname(os.path.abspath(__file__))
_BACKEND = os.path.dirname(_HERE)
sys.path.insert(0, _HERE)
sys.path.insert(0, _BACKEND)

# 必须在导入 app 之前安装 DNS 兜底（未设置 LM_TEST_DNS_MAP 时为空操作）
import dns_shim  # noqa: E402

_DNS_ENTRIES = dns_shim.install()

_DB_PATH = os.path.join(_HERE, "_smoke_s25.db")
if os.path.exists(_DB_PATH):
    os.remove(_DB_PATH)
os.environ["DATABASE_URL"] = "sqlite:///" + _DB_PATH.replace("\\", "/")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from app.core.config import settings  # noqa: E402
from app.core.database import Base, SessionLocal, engine  # noqa: E402
from app.core.security import create_admin_token, create_admin_token_for_user  # noqa: E402
from app.main import app  # noqa: E402
from app.models.models import (  # noqa: E402
    AdminMenu,
    AdminRole,
    AdminRoleMenu,
    AdminUserRole,
    Question,
    User,
)
from app.schemas.schemas import QuestionOut  # noqa: E402
from app.services import llm_service  # noqa: E402

GENERATE_URL = "/api/admin/ai/generate-questions"
BATCH_SAVE_URL = "/api/admin/ai/questions/batch-save"
QUESTIONS_URL = "/api/admin/questions"

RESULTS = []
API_KEY = settings.LLM_API_KEY.strip()
# 密钥指纹：取首尾片段用于断言响应未泄露，断言输出本身不构成完整密钥
KEY_HEAD, KEY_TAIL = API_KEY[:16], API_KEY[-16:]


def record(case_id, name, ok, detail=""):
    RESULTS.append({"id": case_id, "name": name, "ok": bool(ok), "detail": str(detail)[:400]})
    print(("PASS " if ok else "FAIL ") + case_id + " " + name
          + ((" | " + str(detail)[:220]) if detail else ""))


def q(fn):
    db = SessionLocal()
    try:
        return fn(db)
    finally:
        db.close()


def build_xlsx(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "题目数据"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


IMPORT_HEADERS = ["题型", "题目内容", "选项A", "选项B", "选项C", "选项D", "选项E", "选项F",
                  "正确答案", "答案解析", "分值", "知识分类"]


def fake_item(question, labels=("A", "B", "C", "D"), answer="B", analysis="解析"):
    """构造一条 模型风格的题目元素（type 由 _normalize_item 强制为 single）。"""
    return {
        "type": "single",
        "question": question,
        "options": [{"label": l, "text": f"{question}选项{l}"} for l in labels],
        "answer": [answer] if answer else [],
        "analysis": analysis,
    }


def patch_prompt(monkey, reply=None, exc=None):
    """把 llm_service.prompt 替换为固定回复或固定异常。"""
    def _fake(text, timeout=None):
        if exc is not None:
            raise exc
        return reply if isinstance(reply, str) else json.dumps(reply, ensure_ascii=False)
    monkey.append(llm_service.prompt)
    llm_service.prompt = _fake


def restore_prompt(monkey):
    while monkey:
        llm_service.prompt = monkey.pop()


def no_key_leak(*responses):
    """断言响应体不含 API Key 的首尾指纹。"""
    for resp in responses:
        body = resp.text
        if KEY_HEAD in body or KEY_TAIL in body:
            return False
    return True


def draft_payload(items):
    return {"questions": [it["data"] for it in items]}


def main():
    print(f"DNS 映射条目：{_DNS_ENTRIES}（0 表示走系统 DNS）")
    print(f"大模型配置：base_url={'已配置' if settings.LLM_BASE_URL else '缺失'} "
          f"model={settings.LLM_MODEL or '缺失'} "
          f"api_key={'已配置' if API_KEY else '缺失'}")
    Base.metadata.create_all(engine)
    client = TestClient(app)
    admin_h = {"Authorization": "Bearer " + create_admin_token("s25-smoke")}

    # 打桩用例的 stub 对每一批都返回同一份固定列表，批数一变条数就翻倍、顺序断言随之失效。
    # 故打桩阶段固定单批，使离线断言与 AI_QUESTION_BATCH_SIZE 这个调优旋钮解耦；
    # A14/A15 真实调用前恢复，以便真正验证分批行为。
    real_batch_size = settings.AI_QUESTION_BATCH_SIZE
    settings.AI_QUESTION_BATCH_SIZE = 1000

    # 无 questions 菜单的管理员（仅 dashboard），用于 403 断言
    def _seed_restricted_admin(db):
        user = User(openid="s25-restricted-openid", nickname="s25受限管理员", status="active")
        db.add(user)
        db.flush()
        menu = AdminMenu(code="dashboard", name="驾驶舱", sort_order=0)
        role = AdminRole(code="s25viewer", name="仅驾驶舱")
        db.add_all([menu, role])
        db.flush()
        db.add(AdminRoleMenu(role_id=role.id, menu_id=menu.id))
        db.add(AdminUserRole(user_id=user.id, role_id=role.id, enabled=True))
        db.commit()
        return user.id
    restricted_id = q(_seed_restricted_admin)
    restricted_h = {"Authorization": "Bearer " + create_admin_token_for_user(restricted_id)}

    # ---------------- A01 / A02 Excel 导入回归 ----------------
    r_tpl = client.get("/api/admin/questions/import-template", headers=admin_h)
    tpl_ok = False
    tpl_sheets = 0
    if r_tpl.status_code == 200:
        wb = load_workbook(io.BytesIO(r_tpl.content))
        tpl_sheets = len(wb.sheetnames)
        ws = wb["题目数据"] if "题目数据" in wb.sheetnames else wb.active
        tpl_ok = ("spreadsheetml" in r_tpl.headers.get("content-type", "")
                  and tpl_sheets >= 2
                  and [c.value for c in ws[1]] == IMPORT_HEADERS)
    record("A01", "Excel 模板下载未受影响（内容类型/双工作表/表头）", tpl_ok,
           f"status={r_tpl.status_code} ctype={r_tpl.headers.get('content-type', '')[:40]} sheets={tpl_sheets}")

    rows = [["单选", f"Excel 回归题{i}", "甲", "乙", "丙", "丁", "", "", "B", "解析", 20, "event"]
            for i in range(1, 4)]
    files = {"file": ("regression.xlsx", build_xlsx(IMPORT_HEADERS, rows),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r_prev = client.post("/api/admin/questions/import-preview", headers=admin_h, files=files)
    prev = r_prev.json() if r_prev.status_code == 200 else {}
    r_conf = client.post("/api/admin/questions/import-confirm", headers=admin_h,
                         json={"preview_token": prev.get("preview_token", "")})
    baseline_ids = q(lambda db: [x.id for x in db.query(Question).order_by(Question.id).all()])
    record("A02", "Excel 预览 + 确认导入未受影响",
           r_prev.status_code == 200 and prev.get("valid_count") == 3 and prev.get("invalid_count") == 0
           and r_conf.status_code == 200 and r_conf.json().get("imported") == 3
           and len(baseline_ids) == 3 and baseline_ids == sorted(baseline_ids),
           f"preview={prev.get('valid_count')}/{prev.get('invalid_count')} confirm={r_conf.json()} ids={baseline_ids}")

    # ---------------- A03 参数校验 ----------------
    cases = [
        ("空主题", {"topic": "  ", "count": 10}, 400),
        ("题型非单选", {"topic": "长征", "count": 10, "question_type": "judge"}, 400),
        ("难度非法", {"topic": "长征", "count": 10, "difficulty": "extreme"}, 400),
        ("分类非法", {"topic": "长征", "count": 10, "category": "poem"}, 400),
        ("数量为 0", {"topic": "长征", "count": 0}, 422),
        ("数量超上限", {"topic": "长征", "count": 500}, 422),
    ]
    monkey = []
    patch_prompt(monkey, reply=[fake_item("不应被调用")])
    param_detail = []
    param_ok = True
    for name, payload, expect in cases:
        resp = client.post(GENERATE_URL, headers=admin_h, json=payload)
        param_detail.append(f"{name}={resp.status_code}")
        if resp.status_code != expect:
            param_ok = False
    restore_prompt(monkey)
    record("A03", "生成接口参数校验（空主题/题型/难度/分类/数量）", param_ok, " ".join(param_detail))

    # ---------------- A04 鉴权 ----------------
    r_anon = client.post(GENERATE_URL, json={"topic": "长征", "count": 10})
    r_anon_save = client.post(BATCH_SAVE_URL, json={"questions": []})
    r_low = client.post(GENERATE_URL, headers=restricted_h, json={"topic": "长征", "count": 10})
    r_low_save = client.post(BATCH_SAVE_URL, headers=restricted_h,
                             json=draft_payload([{"data": fake_item("越权题"), "index": 1, "errors": []}]))
    record("A04", "未登录 401、无 questions 菜单的管理员 403（生成与保存）",
           r_anon.status_code == 401 and r_anon_save.status_code == 401
           and r_low.status_code == 403 and r_low_save.status_code == 403,
           f"anon={r_anon.status_code}/{r_anon_save.status_code} restricted={r_low.status_code}/{r_low_save.status_code}")

    # ---------------- A05 模型故障映射 ----------------
    faults = [
        ("超时", llm_service.LlmError("AI 生成超时，请减少数量或稍后重试"), 502, "超时"),
        ("鉴权失败", llm_service.LlmError("AI 服务鉴权失败，请检查 LLM_API_KEY 是否有效"), 502, "鉴权"),
        ("服务不可用", llm_service.LlmError("AI 服务暂时不可用，请稍后重试"), 502, "不可用"),
        ("回复截断", llm_service.LlmError("AI 回复未正常结束（max_tokens），请减少单次生成数量后重试"), 502, "未正常结束"),
        ("无效输出", ValueError("AI 返回的 JSON 无法解析，请重试"), 400, "JSON"),
    ]
    fault_detail, fault_ok = [], True
    for name, exc, expect, keyword in faults:
        monkey = []
        patch_prompt(monkey, exc=exc)
        resp = client.post(GENERATE_URL, headers=admin_h, json={"topic": "长征", "count": 10})
        restore_prompt(monkey)
        detail_text = str(resp.json().get("detail", ""))
        hit = resp.status_code == expect and keyword in detail_text
        fault_detail.append(f"{name}={resp.status_code}")
        if not hit:
            fault_ok = False
            fault_detail[-1] += f"(detail={detail_text[:40]})"
        if not no_key_leak(resp):
            fault_ok = False
    record("A05", "模型故障映射为明确中文提示（超时/鉴权/服务异常/截断/无效 JSON）", fault_ok,
           " ".join(fault_detail))

    # ---------------- A06 输出容错 ----------------
    fenced = "好的，以下是题目：\n```json\n" + json.dumps(
        [fake_item("围栏容错题1"), fake_item("围栏容错题2")], ensure_ascii=False) + "\n```\n hope it helps"
    monkey = []
    patch_prompt(monkey, reply=fenced)
    r_fence = client.post(GENERATE_URL, headers=admin_h, json={"topic": "长征", "count": 2})
    restore_prompt(monkey)
    fence = r_fence.json() if r_fence.status_code == 200 else {}
    record("A06", "Markdown 代码围栏与前后多余文本仍能解析出题",
           r_fence.status_code == 200 and fence.get("generated") == 2 and fence.get("valid_count") == 2,
           f"status={r_fence.status_code} generated={fence.get('generated')} valid={fence.get('valid_count')}")

    # ---------------- A07 逐题问题识别 ----------------
    messy = [
        fake_item("合法题A"),
        fake_item("答案不在选项中", answer="Z"),
        fake_item("选项不足", labels=("A",)),
        fake_item("", answer="A"),
        fake_item("Excel 回归题1"),          # 与 A02 已导入题目重复
        fake_item("合法题A"),                 # 批内重复
        {"question": "结构非法题", "options": "not-a-list", "answer": ["A"]},
    ]
    monkey = []
    patch_prompt(monkey, reply=messy)
    r_messy = client.post(GENERATE_URL, headers=admin_h, json={"topic": "长征", "count": 7})
    restore_prompt(monkey)
    messy_out = r_messy.json() if r_messy.status_code == 200 else {}
    items = messy_out.get("items", [])

    def errs_at(i):
        return items[i]["errors"] if i < len(items) else ["<草稿缺失>"]

    # 草稿顺序与 模型输出顺序一致（单批生成），按序号断言，避免同题干在字典里互相覆盖
    struct_ok = (len(items) > 6 and items[6]["data"]["question"] == ""
                 and len(items[6]["errors"]) > 0)
    checks = {
        "顺序一致": len(items) == len(messy),
        "合法题无错": errs_at(0) == [],
        "答案不在选项": any("不在选项标签中" in e for e in errs_at(1)),
        "选项不足": any("至少需要 2 个选项" in e for e in errs_at(2)),
        "题干为空": any("题干不能为空" in e for e in errs_at(3)),
        "与库内重复": any("与题库中已有题目重复" in e for e in errs_at(4)),
        "批内重复": any("与本次生成的其他题目重复" in e for e in errs_at(5)),
        "结构非法占位": struct_ok,
        "统计一致": (messy_out.get("invalid_count") == sum(1 for it in items if it["errors"])
                     and messy_out.get("valid_count") == 1),
    }
    record("A07", "逐题识别答案缺失/选项不完整/题干为空/库内重复/批内重复/结构非法",
           all(checks.values()),
           f"{ {k: v for k, v in checks.items() if not v} or 'all pass' } "
           f"valid={messy_out.get('valid_count')} invalid={messy_out.get('invalid_count')}")

    # ---------------- A08 生成不入库 ----------------
    after_gen_ids = q(lambda db: [x.id for x in db.query(Question).order_by(Question.id).all()])
    record("A08", "生成结果只是草稿，未写入正式题库",
           after_gen_ids == baseline_ids, f"before={baseline_ids} after={after_gen_ids}")

    # ---------------- A09 / A10 / A11 批量保存 ----------------
    good = [fake_item(f"AI 保存题{i}") for i in range(1, 4)]
    bad_answer = fake_item("答案越界题", answer="Q")
    dup_db = fake_item("Excel 回归题2")        # 与库内重复
    dup_self_a = fake_item("批内重复保存题")
    dup_self_b = fake_item("批内重复保存题")
    # 提交顺序：1 合法 / 2 答案越界 / 3 合法 / 4 与库内重复 / 5 批内重复首条(合法) /
    #           6 批内重复次条(拒) / 7 合法 → 预期 saved=4、failed=3
    payload = {"questions": [good[0], bad_answer, good[1], dup_db, dup_self_a, dup_self_b, good[2]]}
    before_save = q(lambda db: db.query(Question).count())
    r_save = client.post(BATCH_SAVE_URL, headers=admin_h, json=payload)
    save_out = r_save.json() if r_save.status_code == 200 else {}
    after_save = q(lambda db: db.query(Question).count())
    # 本批新写入的全部 id（含批内重复首条），应自 max(id)+1 起连续分配
    new_ids = q(lambda db: [x.id for x in db.query(Question)
                            .filter(Question.id > before_save).order_by(Question.id).all()])
    saved_ids = q(lambda db: [x.id for x in db.query(Question)
                              .filter(Question.question.like("AI 保存题%")).order_by(Question.id).all()])
    failures = save_out.get("failures", [])
    fail_reasons = {f["index"]: f["reason"] for f in failures}
    expect_new_ids = list(range(before_save + 1, before_save + 5))
    record("A09", "批量保存重新校验：合法项入库且 id 自 max+1 连续、非法项计入 failures",
           r_save.status_code == 200 and save_out.get("saved") == 4 and save_out.get("failed") == 3
           and after_save == before_save + 4
           and new_ids == expect_new_ids
           and saved_ids == [expect_new_ids[0], expect_new_ids[1], expect_new_ids[3]]
           and "不在选项标签中" in fail_reasons.get(2, "")
           and sorted(fail_reasons) == [2, 4, 6],
           f"saved={save_out.get('saved')} failed={save_out.get('failed')} "
           f"new_ids={new_ids} expect={expect_new_ids} ai_ids={saved_ids} reasons={fail_reasons}")

    record("A10", "批量保存拦截库内重复与批内重复",
           "与题库中已有题目重复" in fail_reasons.get(4, "")
           and fail_reasons.get(5, "") == ""
           and "与本次生成的其他题目重复" in fail_reasons.get(6, ""),
           f"idx4={fail_reasons.get(4)} idx5={fail_reasons.get(5)!r} idx6={fail_reasons.get(6)}")

    r_empty = client.post(BATCH_SAVE_URL, headers=admin_h, json={"questions": []})
    record("A11", "批量保存空列表返回 400",
           r_empty.status_code == 400 and r_empty.json().get("detail") == "没有需要保存的题目",
           f"status={r_empty.status_code} detail={r_empty.json().get('detail')}")

    # ---------------- A12 列表可见且同构 ----------------
    r_list = client.get(QUESTIONS_URL, headers=admin_h, params={"keyword": "AI 保存题", "qtype": "single"})
    listed = r_list.json().get("items", []) if r_list.status_code == 200 else []
    excel_item = q(lambda db: db.query(Question).filter(Question.question == "Excel 回归题1").first())
    ai_item = q(lambda db: db.query(Question).filter(Question.question == "AI 保存题1").first())
    same_shape = excel_item is not None and ai_item is not None and (
        isinstance(ai_item.options, list) and isinstance(ai_item.answer, list)
        and set(ai_item.options[0].keys()) == set(excel_item.options[0].keys())
        and ai_item.type == excel_item.type == "single" and ai_item.active is True)
    record("A12", "保存后题库列表可查到，且与 Excel 导入题目字段同构",
           r_list.status_code == 200 and r_list.json().get("total") == 3 and len(listed) == 3 and same_shape,
           f"total={r_list.json().get('total') if r_list.status_code == 200 else '-'} same_shape={same_shape}")

    # ---------------- A13 密钥不泄露 ----------------
    monkey = []
    patch_prompt(monkey, reply=[fake_item("密钥泄露检查题")])
    r_leak = client.post(GENERATE_URL, headers=admin_h, json={"topic": "长征", "count": 1})
    restore_prompt(monkey)
    r_leak_save = client.post(BATCH_SAVE_URL, headers=admin_h,
                              json={"questions": [fake_item("密钥泄露检查题")]})
    r_leak_list = client.get(QUESTIONS_URL, headers=admin_h)
    record("A13", "生成/保存/列表响应均不泄露 LLM_API_KEY",
           no_key_leak(r_leak, r_leak_save, r_leak_list, r_messy, r_save),
           f"key_configured={bool(API_KEY)} checked=5 responses")

    # ---------------- A16 小程序答题契约 ----------------
    contract_item = q(lambda db: db.query(Question).filter(Question.question == "AI 保存题1").first())
    dumped = QuestionOut.model_validate(contract_item).model_dump(by_alias=True) if contract_item else {}
    record("A16", "小程序下发契约不变：camelCase 且不含 answer/analysis",
           set(dumped.keys()) == {"id", "type", "question", "options", "score"}
           and dumped.get("type") == "single"
           and all(set(o.keys()) == {"label", "text"} for o in dumped.get("options", [])),
           f"keys={sorted(dumped.keys())}")

    # ---------------- A14 / A15 真实调用模型 ----------------
    settings.AI_QUESTION_BATCH_SIZE = real_batch_size  # 恢复 .env 配置，真实验证分批
    if not API_KEY:
        record("A14", "真实生成 10 道单选题", False, "跳过：LLM_API_KEY 未配置")
        record("A15", "真实生成 50 道题（分批长响应）", False, "跳过：LLM_API_KEY 未配置")
    else:
        r10 = client.post(GENERATE_URL, headers=admin_h, json={
            "topic": "红军长征中的重要战役", "count": 10, "question_type": "single",
            "difficulty": "medium", "category": "event",
            "extra_requirements": "史实准确，避免重复，附简明解析"})
        try:
            body10 = r10.json()
        except ValueError:
            body10 = {}
        out10 = body10 if r10.status_code == 200 else {}
        # 非 200 时保留上游错误详情，否则 502 只剩状态码无法定位（截断/超时/鉴权）
        err10 = "" if r10.status_code == 200 else str(body10.get("detail") or r10.text)[:200]
        items10 = out10.get("items", [])
        struct_ok = all(
            it["data"]["type"] == "single"
            and len(it["data"]["options"]) >= 2
            and len({o["label"] for o in it["data"]["options"]}) == len(it["data"]["options"])
            and len(it["data"]["answer"]) == 1
            and it["data"]["answer"][0] in {o["label"] for o in it["data"]["options"]}
            and it["data"]["question"].strip()
            and it["data"]["analysis"].strip()
            for it in items10)
        uniq10 = len({it["data"]["question"].strip() for it in items10}) == len(items10)
        after10 = q(lambda db: db.query(Question).count())
        record("A14", "真实调用模型生成 10 道单选题：结构合规、答案唯一且在选项中、不入库",
               r10.status_code == 200 and out10.get("generated") == 10 and struct_ok and uniq10
               and not out10.get("batch_errors") and after10 == after_save + 1,
               f"status={r10.status_code} generated={out10.get('generated')} "
               f"valid={out10.get('valid_count')} invalid={out10.get('invalid_count')} "
               f"batch_errors={out10.get('batch_errors')} count_after={after10}"
               + (f" err={err10}" if err10 else ""))
        if items10:
            print("  样例题目：" + json.dumps(items10[0]["data"], ensure_ascii=False)[:300])

        if os.getenv("LM_AI_SKIP_LARGE") == "1":
            record("A15", "真实生成 50 道题（分批长响应）", True, "跳过：LM_AI_SKIP_LARGE=1")
        else:
            r50 = client.post(GENERATE_URL, headers=admin_h, json={
                "topic": "红军长征路线与重要会议", "count": 50, "question_type": "single",
                "difficulty": "mixed", "category": "route"})
            out50 = r50.json() if r50.status_code == 200 else {}
            items50 = out50.get("items", [])
            struct50 = all(
                it["data"]["type"] == "single" and len(it["data"]["answer"]) == 1
                and it["data"]["answer"][0] in {o["label"] for o in it["data"]["options"]}
                for it in items50)
            uniq50 = len({it["data"]["question"].strip() for it in items50})
            record("A15", "真实调用模型生成 50 道题：分批处理长响应、结构与去重达标",
                   r50.status_code == 200 and out50.get("generated") == 50 and struct50
                   and not out50.get("batch_errors") and uniq50 >= 45,
                   f"status={r50.status_code} generated={out50.get('generated')} "
                   f"valid={out50.get('valid_count')} invalid={out50.get('invalid_count')} "
                   f"unique_questions={uniq50} batch_errors={out50.get('batch_errors')}")


def teardown():
    engine.dispose()
    try:
        if os.path.exists(_DB_PATH):
            os.remove(_DB_PATH)
    except OSError as exc:
        print(f"清理临时库失败（可忽略）：{exc}")


if __name__ == "__main__":
    try:
        main()
    finally:
        teardown()
    passed = sum(1 for r in RESULTS if r["ok"])
    summary = {"suite": "s25_ai_question_generate", "total": len(RESULTS),
               "passed": passed, "failed": len(RESULTS) - passed, "results": RESULTS}
    out_path = os.path.join(_HERE, "s25-ai-question-generate-results.json")
    with io.open(out_path, "w", encoding="utf-8") as fp:
        json.dump(summary, fp, ensure_ascii=False, indent=2)
    print(f"\n合计 {len(RESULTS)} 项：通过 {passed}，失败 {len(RESULTS) - passed}")
    print("结果已写入 " + out_path)
    sys.exit(0 if passed == len(RESULTS) else 1)
